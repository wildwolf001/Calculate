"""M1/M2: generic CSV/XLSX parsing, relations and semantic drafts."""
from __future__ import annotations

import json

import yaml
from openpyxl import Workbook

from app.datasets.context import load_context
from app.datasets.manifest import compute_data_version
from app.domain.pipeline import run_query
from app.domain.plan import FilterSpec, QueryPlan
from app.domain.semantic.loader import load_semantic
from app.ingest.documentation import apply_documentation, parse_document
from app.ingest.generic_pipeline import analyze_files


def _sem(root):
    return load_semantic(root / "semantic.yaml", root / "binding.yaml",
                         catalog_path=root / "catalog.json")


def _draft(root):
    return json.loads((root / "semantic_draft.json").read_text(encoding="utf-8"))


def _sum_metric(root, fragment):
    draft = _draft(root)
    for metric in draft["metrics"]:
        if fragment in metric["id"]:
            return metric["id"]
    raise AssertionError(f"metric not found: {fragment}")


def _query(root, metric, dimensions=None):
    sem = _sem(root)
    plan = QueryPlan(metric=metric, dimensions=dimensions or [], limit=100)
    return run_query(plan, sem, db_path=root / "warehouse.duckdb")


def test_simple_csv_types_grain_and_query(tmp_path):
    source = tmp_path / "orders.csv"
    source.write_text(
        "order_id,customer_id,order_date,amount,status,zip\n"
        "1001,C001,2025-01-03,120.50,paid,010001\n"
        "1002,C002,2025-02-04,80,paid,020002\n"
        "1003,C001,2025-02-05,200,cancel,010001\n",
        encoding="utf-8",
    )
    root = tmp_path / "ds"
    out = analyze_files([source], root, "ds_simple_csv", "简单销售")
    table = out.tables[0]
    by_label = {c.label: c for c in table.columns}
    assert by_label["order_id"].inferred_type == "integer"
    assert by_label["order_date"].inferred_type == "date"
    assert by_label["amount"].inferred_type == "decimal"
    assert by_label["zip"].inferred_type == "varchar"
    assert by_label["zip"].samples[0] == "010001"
    assert table.grain == ["c01"]
    assert table.duplicate_rows == 0

    semantic_doc = yaml.safe_load((root / "semantic.yaml").read_text(encoding="utf-8"))
    # Data max is 2025-02-05; semantic ranges are half-open, so end is +1 day.
    assert semantic_doc["time_range"] == ["2025-01-03", "2025-02-06"]

    draft = _draft(root)
    assert draft["status"] == "draft"
    assert draft["facts"] == ["t_001"]
    metric = _sum_metric(root, "_c04_sum")
    outcome = _query(root, metric)
    assert outcome.ok, outcome.error
    assert outcome.result.rows[0][0] == 400.5


def test_gbk_semicolon_header_offset_and_totals(tmp_path):
    source = tmp_path / "销售.csv"
    text = (
        "标题;;;\n"
        "生成时间;2025-01-01;;\n"
        "订单号;客户;金额;状态\n"
        "001;C1;10;paid\n"
        "002;C2;20;paid\n"
        "合计;;30;\n"
    )
    source.write_bytes(text.encode("gb18030"))
    root = tmp_path / "ds"
    out = analyze_files([source], root, "ds_gbk", "GBK 销售")
    table = out.tables[0]
    assert table.encoding in {"gb18030", "gbk"}
    assert table.delimiter == ";"
    assert table.header_row == 2
    assert table.rows == 2
    assert table.totals_rows == [2]
    by_label = {c.label: c for c in table.columns}
    assert by_label["订单号"].inferred_type == "varchar"
    metric = _sum_metric(root, "_c03_sum")
    outcome = _query(root, metric)
    assert outcome.ok, outcome.error
    assert outcome.result.rows[0][0] == 30


def test_multi_file_relation_and_joined_dimension(tmp_path):
    orders = tmp_path / "orders.csv"
    customers = tmp_path / "customers.csv"
    orders.write_text(
        "order_id,customer_id,order_date,amount\n"
        "1,C1,2025-01-03,100\n"
        "2,C2,2025-01-04,50\n"
        "3,C1,2025-02-03,25\n",
        encoding="utf-8",
    )
    customers.write_text(
        "customer_id,customer_name,city\n"
        "C1,Alice,Beijing\n"
        "C2,Bob,Shanghai\n",
        encoding="utf-8",
    )
    root = tmp_path / "ds"
    out = analyze_files([orders, customers], root, "ds_multi", "多表")
    draft = _draft(root)
    accepted = [r for r in draft["relations"] if r["accepted"]]
    assert len(accepted) == 1
    assert accepted[0]["child_col"] == "c02"
    assert accepted[0]["parent_col"] == "c01"
    dims = [d for d in draft["dimensions"]
            if d.get("join_path") == ["t_001 -> t_002"]]
    assert dims
    city_dim = next((d for d in dims if d.get("name_zh") == "城市"), dims[0])
    metric = _sum_metric(root, "_c04_sum")
    outcome = _query(root, metric, dimensions=[city_dim["id"]])
    assert outcome.ok, outcome.error
    rows = {row[0]: row[1] for row in outcome.result.rows}
    assert rows == {"Beijing": 125, "Shanghai": 50}


def test_xlsx_multiple_sheets(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售"
    sheet.append(["日期", "金额", "状态"])
    sheet.append(["2025-01-01", 10, "paid"])
    sheet.append(["2025-01-02", 20, "paid"])
    other = workbook.create_sheet("客户")
    other.append(["客户", "城市"])
    other.append(["C1", "北京"])
    other.append(["C2", "上海"])
    workbook.save(path)

    root = tmp_path / "ds"
    out = analyze_files([path], root, "ds_xlsx", "Excel 多表")
    assert len(out.tables) == 2
    labels = " ".join(t.label for t in out.tables)
    assert "销售" in labels and "客户" in labels
    assert all(t.rows == 2 for t in out.tables)


def test_relation_dimension_filters_use_display_value(tmp_path):
    orders = tmp_path / "orders.csv"
    customers = tmp_path / "customers.csv"
    orders.write_text(
        "order_id,customer_id,order_date,amount\n"
        "1,C1,2025-01-03,100\n"
        "2,C2,2025-01-04,50\n",
        encoding="utf-8",
    )
    customers.write_text(
        "customer_id,customer_name,city\n"
        "C1,张三,北京\n"
        "C2,李四,上海\n",
        encoding="utf-8",
    )
    root = tmp_path / "ds"
    analyze_files([orders, customers], root, "ds_filter", "筛选")
    sem = _sem(root)
    draft = _draft(root)
    ids = {t["label"]: t["table_id"] for t in draft["tables"]}
    city = next(d for d in sem.dimensions.values()
                if d.raw.get("name_zh") == "城市")
    assert city.raw["filter_column"] == city.raw["column"]
    metric = next(
        m for m in sem.metrics.values()
        if m.raw.get("fact") == ids["orders"] and "SUM(" in str(m.raw.get("expr"))
    )
    plan = QueryPlan(
        metric=metric.id,
        filters={city.id: FilterSpec(op="eq", value="北京")},
        limit=10,
    )
    outcome = run_query(plan, sem, db_path=root / "warehouse.duckdb")
    assert outcome.ok, outcome.error
    assert outcome.result.rows[0][0] == 100


def test_high_cardinality_label_is_a_dimension(tmp_path):
    products = tmp_path / "products.csv"
    items = tmp_path / "order_items.csv"
    products.write_text(
        "product_id,product_name\n"
        + "".join(f"P{i:03d},商品{i:03d}\n" for i in range(250)),
        encoding="utf-8",
    )
    items.write_text(
        "order_id,product_id,price\n"
        + "".join(f"{i},P{i % 250:03d},10\n" for i in range(300)),
        encoding="utf-8",
    )
    root = tmp_path / "ds"
    analyze_files([items, products], root, "ds_high_card", "高基数")
    sem = _sem(root)
    draft = _draft(root)
    ids = {t["label"]: t["table_id"] for t in draft["tables"]}
    cols = {d.raw.get("column") for d in sem.dimensions.values()}
    assert f"{ids['products']}.c02" in cols  # product_name


def test_doc_entity_rule_with_prefixed_sheet_labels(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    orders = workbook.active
    orders.title = "orders"
    orders.append(["order_id", "customer_id", "order_status", "purchase_at"])
    orders.append(["O1", "C1", "delivered", "2025-01-01"])
    orders.append(["O2", "C2", "delivered", "2025-01-02"])
    orders.append(["O3", "C1", "cancelled", "2025-01-03"])
    orders.append(["O4", "C1", "delivered", "2025-01-04"])
    customers = workbook.create_sheet("customers")
    customers.append(["customer_id", "customer_unique_id"])
    customers.append(["C1", "U1"])
    customers.append(["C2", "U2"])
    workbook.save(path)

    root = tmp_path / "ds"
    analyze_files([path], root, "ds_entity", "实体")
    ctx = load_context(
        dataset_id="ds_entity", name="实体", root=root,
        warehouse=root / "warehouse.duckdb",
        semantic_path=root / "semantic.yaml",
        binding_path=root / "binding.yaml", mode="draft",
        data_version=compute_data_version(root / "warehouse.duckdb"),
        catalog_path=root / "catalog.json",
    )
    doc = parse_document(
        "复购客户：同一 customer_unique_id 完成至少 2 笔已送达订单。",
        "metrics.md", ctx,
    )
    result = apply_documentation(ctx, doc)
    assert result["applied"], result["warnings"]
    sem = _sem(root)
    repeat = next(m for m in sem.metrics.values() if m.name == "复购客户")
    outcome = run_query(QueryPlan(metric=repeat.id, limit=5), sem,
                        db_path=root / "warehouse.duckdb")
    assert outcome.ok, outcome.error
    assert outcome.result.rows[0][0] == 1


def test_only_one_selected_metric_per_name(tmp_path):
    orders = tmp_path / "orders.csv"
    items = tmp_path / "order_items.csv"
    customers = tmp_path / "customers.csv"
    orders.write_text(
        "order_id,customer_id,order_date\n"
        "1,C1,2025-01-01\n2,C2,2025-01-02\n",
        encoding="utf-8",
    )
    items.write_text(
        "order_id,order_item_id,price\n"
        "1,1,10\n1,2,20\n2,1,5\n",
        encoding="utf-8",
    )
    customers.write_text(
        "customer_id,customer_name\nC1,张三\nC2,李四\n",
        encoding="utf-8",
    )
    root = tmp_path / "ds"
    analyze_files([orders, items, customers], root, "ds_dupe", "去重")
    sem = _sem(root)
    selected: dict[str, str] = {}
    for metric in sem.metrics.values():
        if metric.raw.get("selected", True):
            assert metric.name not in selected, (metric.name, selected[metric.name], metric.id)
            selected[metric.name] = metric.id


def test_no_header_and_duplicate_rows_are_flagged(tmp_path):
    source = tmp_path / "raw.csv"
    source.write_text("1001,C1,10\n1002,C2,20\n1002,C2,20\n", encoding="utf-8")
    root = tmp_path / "ds"
    out = analyze_files([source], root, "ds_nohdr", "无表头")
    table = out.tables[0]
    assert table.header_row is None
    assert table.columns[0].label.startswith("column_")
    assert table.grain == []
    assert table.duplicate_rows == 1
    quality = json.loads((root / "quality_report.json").read_text(encoding="utf-8"))
    ids = {c["id"] for c in quality["checks"]}
    assert "duplicate_full_rows" in ids
    assert "no_unique_key" in ids
