"""File adapters: CSV/TSV and XLSX/XLSM -> in-memory raw rows (all strings).

The adapters never coerce types. They only decode, split and stringify; type
inference happens later and is allowed to fall back to VARCHAR.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_ROWS = 1_000_000
MAX_CELLS = 20_000_000

TEXT_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass
class Unit:
    file: Path
    sheet: str | None
    label: str
    hidden: bool = False


@dataclass
class RawTable:
    unit: Unit
    rows: list[list[str]] = field(default_factory=list)
    width: int = 0
    warnings: list[str] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None
    truncated: bool = False


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    return str(value).strip()


def _decode_bytes(data: bytes) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig", warnings
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16", warnings
    try:
        return data.decode("utf-8"), "utf-8", warnings
    except UnicodeDecodeError:
        pass
    for encoding in ("gb18030", "big5"):
        try:
            return data.decode(encoding), encoding, warnings
        except UnicodeDecodeError:
            continue
    warnings.append("无法确定编码，按 latin-1 读取，中文可能损坏")
    return data.decode("latin-1"), "latin-1", warnings


def _sniff_delimiter(text: str, suffix: str) -> str:
    if suffix == ".tsv":
        return "\t"
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        pass
    lines = [line for line in sample.splitlines() if line.strip()][:20]
    best, best_score = ",", 0.0
    for delim in [",", ";", "\t", "|"]:
        counts = [line.count(delim) for line in lines]
        if not counts:
            continue
        non_zero = [c for c in counts if c > 0]
        if not non_zero:
            continue
        consistency = len(non_zero) / len(counts)
        score = consistency * (sum(counts) / len(counts))
        if score > best_score:
            best, best_score = delim, score
    return best


def open_workbook(path: Path):
    """Open an Excel workbook once; callers own the returned handle."""
    from openpyxl import load_workbook
    return load_workbook(path, read_only=True, data_only=True, keep_vba=False)


def discover_units(files: list[Path],
                   workbook_for: Callable[[Path], Any] | None = None) -> list[Unit]:
    """List sheets/units. A shared workbook cache avoids re-parsing the file."""
    units: list[Unit] = []
    for path in files:
        suffix = path.suffix.lower()
        if suffix in TEXT_SUFFIXES:
            units.append(Unit(file=path, sheet=None, label=path.stem))
        elif suffix in EXCEL_SUFFIXES:
            if workbook_for is None:
                workbook = open_workbook(path)
            else:
                workbook = workbook_for(path)
            try:
                for name in workbook.sheetnames:
                    sheet = workbook[name]
                    state = getattr(sheet, "sheet_state", "visible")
                    units.append(Unit(
                        file=path, sheet=name, label=f"{path.stem}/{name}",
                        hidden=state in {"hidden", "veryHidden"},
                    ))
            finally:
                if workbook_for is None:
                    workbook.close()
        else:
            raise ValueError(f"不支持的文件类型: {path.name}")
    return units


def _read_csv(unit: Unit) -> RawTable:
    path = unit.file
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"文件超过大小上限 {MAX_FILE_BYTES // (1024 * 1024)}MB: {path.name}")
    text, encoding, warnings = _decode_bytes(path.read_bytes())
    delimiter = _sniff_delimiter(text, path.suffix.lower())
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows: list[list[str]] = []
    ragged = 0
    width = 0
    truncated = False
    for idx, row in enumerate(reader):
        if idx >= MAX_ROWS:
            truncated = True
            warnings.append(f"超过行数上限 {MAX_ROWS}，仅分析前 {MAX_ROWS} 行")
            break
        cells = [cell.strip() for cell in row]
        width = max(width, len(cells))
        rows.append(cells)
    for row in rows:
        if len(row) < width:
            ragged += 1
            row.extend([""] * (width - len(row)))
    if ragged:
        warnings.append(f"{ragged} 行短于最宽行，已用空值补齐")
    return RawTable(unit=unit, rows=rows, width=width, warnings=warnings,
                    encoding=encoding, delimiter=delimiter, truncated=truncated)


def _read_xlsx(unit: Unit, workbook=None) -> RawTable:
    path = unit.file
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"文件超过大小上限 {MAX_FILE_BYTES // (1024 * 1024)}MB: {path.name}")
    own_workbook = workbook is None
    if own_workbook:
        workbook = open_workbook(path)
    warnings: list[str] = []
    rows: list[list[str]] = []
    width = 0
    truncated = False
    try:
        sheet = workbook[unit.sheet]
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx >= MAX_ROWS:
                truncated = True
                warnings.append(f"工作表超过行数上限 {MAX_ROWS}，仅分析前 {MAX_ROWS} 行")
                break
            cells = [_cell_text(value) for value in row]
            while cells and cells[-1] == "":
                cells.pop()
            width = max(width, len(cells))
            rows.append(cells)
    finally:
        if own_workbook:
            workbook.close()
    for row in rows:
        if len(row) < width:
            row.extend([""] * (width - len(row)))
    if truncated:
        warnings.append("WARNING_TRUNCATED")
    return RawTable(unit=unit, rows=rows, width=width, warnings=warnings,
                    truncated=truncated)


def read_unit(unit: Unit, workbook=None) -> RawTable:
    suffix = unit.file.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        table = _read_csv(unit)
    elif suffix in EXCEL_SUFFIXES:
        table = _read_xlsx(unit, workbook=workbook)
    else:  # pragma: no cover - discover_units already rejects
        raise ValueError(f"不支持的文件类型: {unit.file.name}")
    cells = sum(len(row) for row in table.rows)
    if cells > MAX_CELLS:
        raise ValueError(
            f"工作表单元格数 {cells} 超过上限 {MAX_CELLS}，请先拆分文件"
        )
    if not table.rows:
        table.warnings.append("工作表为空")
    return table
